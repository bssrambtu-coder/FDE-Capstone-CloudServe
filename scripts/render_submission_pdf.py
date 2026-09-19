"""Render the two net-new submission DOCX files to polished PDFs.

This fallback is used because the packaged LibreOffice renderer is unavailable
and Microsoft Word automation can stall on this host. It preserves document
order, headings, paragraphs, tables and explicit page breaks.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import BaseDocTemplate, Frame, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle

NAVY = colors.HexColor("#17182B")
ORANGE = colors.HexColor("#F06400")
GRID = colors.HexColor("#D9D9D9")


def esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")


def draw_page(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#555965"))
    canvas.drawString(.7 * inch, .38 * inch, "Shashidhar B S  |  CloudServe Solutions")
    canvas.drawRightString(A4[0] - .7 * inch, .38 * inch, f"Page {doc.page}")
    canvas.restoreState()


def styles():
    ss = getSampleStyleSheet()
    return {
        "normal": ParagraphStyle("Body", parent=ss["BodyText"], fontName="Helvetica", fontSize=9.2,
                                 leading=12.2, spaceAfter=7, textColor=colors.black),
        "title": ParagraphStyle("Title", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=25,
                                leading=29, alignment=TA_CENTER, textColor=colors.black, spaceAfter=18),
        "h1": ParagraphStyle("H1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=17,
                             leading=21, textColor=colors.black, spaceBefore=4, spaceAfter=12),
        "h2": ParagraphStyle("H2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=12,
                             leading=15, textColor=colors.black, spaceBefore=9, spaceAfter=6),
        "h3": ParagraphStyle("H3", parent=ss["Heading3"], fontName="Helvetica-Bold", fontSize=10,
                             leading=13, textColor=colors.black, spaceBefore=7, spaceAfter=4),
        "center": ParagraphStyle("Center", parent=ss["BodyText"], fontName="Helvetica", fontSize=11,
                                 leading=15, alignment=TA_CENTER, textColor=ORANGE, spaceAfter=10),
        "cell": ParagraphStyle("Cell", parent=ss["BodyText"], fontName="Helvetica", fontSize=7.1,
                               leading=8.6, textColor=colors.black),
        "head": ParagraphStyle("Head", parent=ss["BodyText"], fontName="Helvetica-Bold", fontSize=7.1,
                               leading=8.6, textColor=colors.white),
    }


def render(src: Path, dst: Path):
    word = Document(src)
    st = styles()
    frame = Frame(.65 * inch, .58 * inch, A4[0] - 1.3 * inch, A4[1] - 1.15 * inch, id="body")
    pdf = BaseDocTemplate(str(dst), pagesize=A4, leftMargin=.65*inch, rightMargin=.65*inch,
                          topMargin=.58*inch, bottomMargin=.58*inch,
                          title=src.stem, author="Shashidhar B S")
    pdf.addPageTemplates(PageTemplate(id="main", frames=[frame], onPage=draw_page))
    story = []
    body = word.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            p = next((x for x in word.paragraphs if x._p is child), None)
            if p is None:
                continue
            has_break = bool(p._p.xpath('.//w:br[@w:type="page"]'))
            if has_break:
                story.append(PageBreak())
            text = p.text.strip()
            if not text:
                continue
            style_name = p.style.name if p.style else "Normal"
            key = "title" if style_name == "Title" else "h1" if style_name == "Heading 1" else "h2" if style_name == "Heading 2" else "h3" if style_name == "Heading 3" else "center" if p.alignment == 1 else "normal"
            story.append(Paragraph(esc(text), st[key]))
        elif child.tag == qn("w:tbl"):
            table = next((x for x in word.tables if x._tbl is child), None)
            if table is None or not table.rows:
                continue
            rows=[]
            for ri,row in enumerate(table.rows):
                rows.append([Paragraph(esc(cell.text.strip()), st["head" if ri == 0 else "cell"]) for cell in row.cells])
            count=len(rows[0]); avail=A4[0]-1.3*inch
            widths=[avail/count]*count
            if count >= 4:
                widths[0]=avail*.15; remaining=avail-widths[0]; widths[1:]=[remaining/(count-1)]*(count-1)
            t=Table(rows,colWidths=widths,repeatRows=1,hAlign="LEFT")
            commands=[("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),
                      ("GRID",(0,0),(-1,-1),.45,GRID),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
                      ("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),
                      ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]
            for ri in range(1,len(rows)):
                if ri%2==0: commands.append(("BACKGROUND",(0,ri),(-1,ri),colors.HexColor("#F3F4F7")))
            t.setStyle(TableStyle(commands)); story.extend([Spacer(1,5),t,Spacer(1,10)])
    pdf.build(story)
    print(dst)


if __name__ == "__main__":
    for name in sys.argv[1:]:
        src=Path(name); render(src,src.with_suffix(".pdf"))
